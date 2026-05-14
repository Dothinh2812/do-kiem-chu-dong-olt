import asyncio
import json
import sys
from pathlib import Path
from datetime import datetime

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

try:
    from do_chu_dong_api.openzca_adapter import OpenZcaClient
    from do_chu_dong_api import notification_service
except ModuleNotFoundError:
    from openzca_adapter import OpenZcaClient
    import notification_service


class DummyClient:
    def __init__(self):
        self.calls = []

    def send_text(self, thread_id, message, group=True):
        self.calls.append(
            {
                "thread_id": thread_id,
                "message": message,
                "group": group,
            }
        )
        return "ok"


def test_openzca_client_builds_command_with_absolute_node_and_profile():
    client = OpenZcaClient(
        binary_path="/opt/openzca/bin/openzca",
        profile="zalo2",
    )

    command = client._build_command(["msg", "send", "123", "hello"])

    assert command[:4] == [
        client.node_path,
        "/opt/openzca/bin/openzca",
        "--profile",
        "zalo2",
    ]
    assert command[4:] == ["msg", "send", "123", "hello"]


def test_send_zalo_message_to_thread_uses_openzca_client():
    client = DummyClient()

    sent = asyncio.run(
        notification_service.send_zalo_message_to_thread(
            "Noi dung canh bao",
            "4761925886931896176",
            client=client,
        )
    )

    assert sent is True
    assert client.calls == [
        {
            "thread_id": "4761925886931896176",
            "message": "Noi dung canh bao",
            "group": True,
        }
    ]


def test_send_zalo_message_to_thread_returns_false_on_transport_error():
    class FailingClient:
        def send_text(self, thread_id, message, group=True):
            raise RuntimeError("openzca unavailable")

    sent = asyncio.run(
        notification_service.send_zalo_message_to_thread(
            "Noi dung canh bao",
            "4761925886931896176",
            client=FailingClient(),
        )
    )

    assert sent is False


def test_send_zalo_message_to_thread_detailed_returns_error_on_transport_error():
    class FailingClient:
        def send_text(self, thread_id, message, group=True):
            raise RuntimeError("openzca unavailable")

    result = asyncio.run(
        notification_service.send_zalo_message_to_thread_detailed(
            "Noi dung canh bao",
            "4761925886931896176",
            client=FailingClient(),
        )
    )

    assert result["success"] is False
    assert result["thread_id"] == "4761925886931896176"
    assert "openzca unavailable" in result["error"]


def test_send_zalo_message_to_thread_detailed_prefers_client_detailed_result():
    class DetailedClient:
        def send_text_detailed(self, thread_id, message, group=True):
            return {
                "success": False,
                "thread_id": thread_id,
                "message": message,
                "error": "stderr message",
                "stdout": "stdout message",
                "stderr": "stderr message",
                "returncode": 7,
                "command": ["/node", "/openzca", "--profile", "zalo2", "msg", "send"],
            }

    result = asyncio.run(
        notification_service.send_zalo_message_to_thread_detailed(
            "Noi dung canh bao",
            "4761925886931896176",
            client=DetailedClient(),
        )
    )

    assert result["success"] is False
    assert result["stderr"] == "stderr message"
    assert result["stdout"] == "stdout message"
    assert result["returncode"] == 7
    assert result["command"] == ["/node", "/openzca", "--profile", "zalo2", "msg", "send"]


def test_send_zalo_message_to_user_detailed_uses_direct_message():
    class DetailedClient:
        def __init__(self):
            self.calls = []

        def send_text_detailed(self, thread_id, message, group=True):
            self.calls.append({"thread_id": thread_id, "message": message, "group": group})
            return {
                "success": True,
                "thread_id": thread_id,
                "message": message,
                "stdout": "ok",
                "stderr": "",
                "returncode": 0,
                "command": ["/node", "/openzca", "msg", "send", thread_id],
            }

    client = DetailedClient()

    result = asyncio.run(
        notification_service.send_zalo_message_to_user_detailed(
            "Noi dung canh bao",
            "user-123",
            client=client,
        )
    )

    assert result["success"] is True
    assert client.calls == [
        {
            "thread_id": "user-123",
            "message": "Noi dung canh bao",
            "group": False,
        }
    ]


def test_send_zalo_message_to_thread_returns_false_without_thread():
    sent = asyncio.run(notification_service.send_zalo_message_to_thread("Noi dung", ""))
    assert sent is False


def test_get_zalo_thread_by_doi_vt_supports_short_aliases():
    assert notification_service.get_zalo_thread_by_doi_vt("Sơn Tây") == "4761925886931896176"
    assert notification_service.get_zalo_thread_by_doi_vt("Quảng Oai") == "7968537750365285360"
    assert notification_service.get_zalo_thread_by_doi_vt("Suối hai") == "6052111621047664"
    assert notification_service.get_zalo_thread_by_doi_vt("Phúc Thọ") == "3142012656522650111"


def test_load_individual_zalo_mapping_from_config_file(tmp_path, monkeypatch):
    mapping_file = tmp_path / "individual_zalo_mapping.json"
    mapping_file.write_text(
        json.dumps(
            {
                "VNPT - Nguyen Van A": "user-a",
                "Nguyen Van B": "user-b",
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    monkeypatch.setenv("INDIVIDUAL_ZALO_MAPPING_FILE", str(mapping_file))

    assert notification_service.get_zalo_user_by_nvkt("VNPT - Nguyen Van A") == "user-a"
    assert notification_service.get_zalo_user_by_nvkt("Nguyen Van B") == "user-b"
    assert notification_service.get_zalo_user_by_nvkt("VNPT - Nguyen Van B") == "user-b"
    assert notification_service.get_zalo_user_by_nvkt("Unknown") is None


def test_get_alert_policy_status_returns_disabled_when_type_is_turned_off():
    policy = notification_service.get_alert_policy_status(
        "outage",
        {
            "enable_individual_alert_notifications": False,
            "individual_alert_time_window": "06:00-18:00",
        },
        now=datetime(2026, 4, 25, 10, 0, 0),
    )

    assert policy["allowed"] is False
    assert policy["reason"] == "disabled"


def test_get_alert_policy_status_respects_time_window():
    policy = notification_service.get_alert_policy_status(
        "wide_area",
        {
            "enable_wide_area_alert_notifications": True,
            "wide_area_alert_time_window": "06:00-18:00",
        },
        now=datetime(2026, 4, 25, 22, 0, 0),
    )

    assert policy["allowed"] is False
    assert policy["reason"] == "outside_time_window"


def test_get_alert_policy_status_supports_group_outage_policy():
    policy = notification_service.get_alert_policy_status(
        "group_outage",
        {
            "enable_group_alert_notifications": False,
            "group_alert_time_window": "06:00-21:00",
        },
        now=datetime(2026, 5, 14, 10, 0, 0),
    )

    assert policy["allowed"] is False
    assert policy["reason"] == "disabled"


def test_group_and_individual_alert_intervals_are_independent():
    config = {
        "group_alert_send_every_batches": 2,
        "individual_alert_send_every_batches": 5,
    }

    assert notification_service.get_group_alert_send_every_batches(config) == 2
    assert notification_service.get_individual_alert_send_every_batches(config) == 5


def test_filter_current_off_alerts_can_use_independent_start_key():
    alerts = [
        {"ma_tb": "GROUP_OLD", "first_off_time": "2026-05-14T06:30:00"},
        {"ma_tb": "GROUP_KEEP", "first_off_time": "2026-05-14T07:30:00"},
    ]

    filtered = notification_service.filter_current_off_alerts_by_cutoff(
        alerts,
        {"group_alert_start_time": "07:00", "individual_alert_start_time": "06:00"},
        start_time_key="group_alert_start_time",
        now=datetime(2026, 5, 14, 8, 0, 0),
    )

    assert [alert["ma_tb"] for alert in filtered] == ["GROUP_KEEP"]


def test_get_alert_policy_status_allows_overnight_window():
    policy = notification_service.get_alert_policy_status(
        "wide_area",
        {
            "enable_wide_area_alert_notifications": True,
            "wide_area_alert_time_window": "22:00-06:00",
        },
        now=datetime(2026, 4, 25, 23, 30, 0),
    )

    assert policy["allowed"] is True
    assert policy["reason"] == "allowed"


def test_filter_current_off_alerts_keeps_only_rows_after_today_cutoff():
    alerts = [
        {"ma_tb": "OLD", "first_off_time": "2026-04-26T05:59:00"},
        {"ma_tb": "KEEP", "first_off_time": "2026-04-26T06:00:00"},
        {"ma_tb": "KEEP2", "first_off_time": "2026-04-26T06:10:00"},
        {"ma_tb": "YESTERDAY", "first_off_time": "2026-04-25T23:00:00"},
        {"ma_tb": "MISSING", "first_off_time": ""},
    ]

    filtered = notification_service.filter_current_off_alerts_by_cutoff(
        alerts,
        {"current_off_alert_start_time": "06:00"},
        now=datetime(2026, 4, 26, 7, 30, 0),
    )

    assert [alert["ma_tb"] for alert in filtered] == ["KEEP", "KEEP2"]


def test_filter_current_off_alerts_allows_empty_cutoff_to_disable_filter():
    alerts = [
        {"ma_tb": "OLD", "first_off_time": "2026-04-25T23:00:00"},
        {"ma_tb": "MISSING", "first_off_time": ""},
    ]

    filtered = notification_service.filter_current_off_alerts_by_cutoff(
        alerts,
        {"current_off_alert_start_time": ""},
        now=datetime(2026, 4, 26, 7, 30, 0),
    )

    assert filtered == alerts


def test_parse_wide_area_excluded_ports_supports_semicolon_list():
    excluded_ports = notification_service.parse_wide_area_excluded_ports(
        "OLT: STY.G22, Port: 0-1-13; OLT: STY.G23, Port: 0-1-14"
    )

    assert ("STY.G22", "0-1-13") in excluded_ports
    assert ("STY.G23", "0-1-14") in excluded_ports


def test_is_wide_area_alert_excluded_matches_short_olt_name_from_mapping(tmp_path, monkeypatch):
    mapping_file = tmp_path / "olt_mapping.xlsx"

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["OLT", "TEN_DSLAM"])
    sheet.append(["HNI.STY.G22.OLT", "STY.G22"])
    workbook.save(mapping_file)

    monkeypatch.setattr(notification_service, "OLT_MAPPING_FILE", str(mapping_file))
    monkeypatch.setattr(notification_service, "_OLT_DISPLAY_NAME_CACHE", None)

    excluded = notification_service.is_wide_area_alert_excluded(
        {
            "olt_name": "HNI.STY.G22.OLT",
            "port": "0-1-13",
        },
        {
            "wide_area_alert_excluded_ports": "OLT: STY.G22, Port: 0-1-13",
        },
    )

    assert excluded is True


def test_append_notification_delivery_log_writes_jsonl_with_preview(tmp_path):
    log_file = tmp_path / "log_message" / "notification_delivery.jsonl"

    notification_service.append_notification_delivery_log(
        {
            "batch_id": "b1",
            "channel": "telegram",
            "alert_type": "outage",
            "status": "SUCCESS",
            "message_full": "Dong 1\nDong 2",
        },
        log_file=str(log_file),
    )

    lines = log_file.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 1
    payload = json.loads(lines[0])
    assert payload["batch_id"] == "b1"
    assert payload["channel"] == "telegram"
    assert payload["alert_type"] == "outage"
    assert payload["status"] == "SUCCESS"
    assert payload["message_full"] == "Dong 1\nDong 2"
    assert payload["message_preview"] == "Dong 1 Dong 2"
    assert "timestamp" in payload


def test_get_olt_display_name_uses_short_name_from_mapping_file(tmp_path, monkeypatch):
    mapping_file = tmp_path / "olt_mapping.xlsx"

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["OLT", "TEN_DSLAM"])
    sheet.append(["HNI.STY.DGM.OLT.ZT.1.1", "DGM.G41"])
    workbook.save(mapping_file)

    monkeypatch.setattr(notification_service, "OLT_MAPPING_FILE", str(mapping_file))
    monkeypatch.setattr(notification_service, "_OLT_DISPLAY_NAME_CACHE", None)

    assert notification_service.get_olt_display_name("HNI.STY.DGM.OLT.ZT.1.1") == "DGM.G41"
    assert notification_service.get_olt_display_name("UNKNOWN.OLT") == "UNKNOWN.OLT"


def test_format_wide_area_outage_message_uses_short_olt_name_from_mapping(tmp_path, monkeypatch):
    mapping_file = tmp_path / "olt_mapping.xlsx"

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["OLT", "TEN_DSLAM"])
    sheet.append(["HNI.STY.DGM.OLT.ZT.1.1", "DGM.G41"])
    workbook.save(mapping_file)

    monkeypatch.setattr(notification_service, "OLT_MAPPING_FILE", str(mapping_file))
    monkeypatch.setattr(notification_service, "_OLT_DISPLAY_NAME_CACHE", None)

    message = notification_service.format_wide_area_outage_message(
        [
            {
                "port": "1/1/1",
                "subscriber_count": 3,
                "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
                "subscriber_list": [],
            }
        ]
    )

    assert "OLT: DGM.G41" in message
    assert "HNI.STY.DGM.OLT.ZT.1.1" not in message


def test_format_wide_area_outage_message_includes_duration_minutes():
    message = notification_service.format_wide_area_outage_message(
        [
            {
                "port": "1/1/1",
                "subscriber_count": 3,
                "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
                "off_duration_minutes": 20,
                "subscriber_list": [],
            }
        ]
    )

    assert "Kéo dài: 20 phút" in message


def test_format_wide_area_outage_message_uses_port_down_title_for_port_down_incident():
    message = notification_service.format_wide_area_outage_message(
        [
            {
                "port": "1/1/1",
                "subscriber_count": 3,
                "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
                "incident_type": "port_down",
                "subscriber_list": [],
            }
        ]
    )

    assert "CẢNH BÁO PORT OLT DOWN" in message
    assert "CẢNH BÁO SỰ CỐ DIỆN RỘNG" not in message


def test_format_consolidated_outage_by_nvkt_includes_duration_minutes():
    message = notification_service.format_consolidated_outage_by_nvkt(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB",
                "dienthoai_lh": "0912345678",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": datetime(2026, 4, 25, 10, 0, 0),
                "off_duration_minutes": 15,
            }
        ]
    )

    assert "TB001 | Ten TB | 0912345678 | 10:00 | 15 phút" in message


def test_format_consolidated_outage_by_nvkt_falls_back_to_first_off_time_delta(monkeypatch):
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 4, 25, 10, 20, 0)

    monkeypatch.setattr(notification_service, "datetime", FrozenDateTime)

    message = notification_service.format_consolidated_outage_by_nvkt(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB",
                "dienthoai_lh": "0912345678",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": datetime(2026, 4, 25, 10, 0, 0),
            }
        ]
    )

    assert "TB001 | Ten TB | 0912345678 | 10:00 | 20 phút" in message


def test_format_wide_area_outage_for_zalo_includes_duration_minutes():
    message = notification_service.format_wide_area_outage_for_zalo(
        {
            "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
            "port": "1/1/1",
            "subscriber_count": 3,
            "off_duration_minutes": 20,
            "subscriber_list": [],
        }
    )

    assert "Kéo dài: 20 phút" in message


def test_format_wide_area_outage_for_zalo_uses_port_down_title_for_port_down_incident():
    message = notification_service.format_wide_area_outage_for_zalo(
        {
            "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
            "port": "1/1/1",
            "subscriber_count": 3,
            "incident_type": "port_down",
            "subscriber_list": [],
        }
    )

    assert "CẢNH BÁO PORT OLT DOWN" in message
    assert "CẢNH BÁO SỰ CỐ DIỆN RỘNG" not in message


def test_format_wide_area_outage_for_zalo_includes_start_time_and_fallback_duration(monkeypatch):
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 4, 27, 10, 0, 0)

    monkeypatch.setattr(notification_service, "datetime", FrozenDateTime)

    message = notification_service.format_wide_area_outage_for_zalo(
        {
            "olt_name": "HNI.STY.DGM.OLT.ZT.1.1",
            "port": "1-1-8",
            "subscriber_count": 20,
            "first_off_time": "2026-04-27T09:40:00",
            "subscriber_list": [],
        }
    )

    assert "Bắt đầu: 27/04/2026 09:40" in message
    assert "Kéo dài: 20 phút" in message


def test_format_consolidated_outage_for_doi_includes_duration_minutes():
    message = notification_service.format_consolidated_outage_for_doi(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB",
                "dienthoai_lh": "0912345678",
                "diachi_lapdat": "123 Duong Rat Dai, Phuong Trung Tam, Thi Xa Son Tay",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:1",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "off_duration_minutes": 15,
                "first_off_time": "2026-04-25T10:30:00",
            }
        ],
        "Sơn Tây",
    )

    assert "Kéo dài: 0,2 giờ" in message
    assert "[TB001] Ten TB - 0912345678" in message
    assert "👷 Nguyen Van A (1 TB)" in message
    assert "OFF: 25/04/2026 10:30" in message
    assert "Đ/c: 123 Duong Rat Dai, Phuong Trung Tam, Thi" in message
    assert "Port: STY.G51_1/1" in message
    assert "Đội: Sơn Tây" not in message


def test_format_consolidated_outage_for_doi_groups_alerts_by_nvkt():
    message = notification_service.format_consolidated_outage_for_doi(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB 1",
                "dienthoai_lh": "0912345678",
                "diachi_lapdat": "Dia chi so 1, phuong A, son tay",
                "port_id": "HNI.STY.DGM.OLT.ZT.1.2_1-1-8:17",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "off_duration_minutes": 15,
            },
            {
                "ma_tb": "TB002",
                "ten_tb": "Ten TB 2",
                "dienthoai_lh": "0987654321",
                "diachi_lapdat": "Dia chi so 2, phuong B, son tay",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-13:2",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "off_duration_minutes": 18,
            },
            {
                "ma_tb": "TB003",
                "ten_tb": "Ten TB 3",
                "dienthoai_lh": "0900000000",
                "diachi_lapdat": "Dia chi so 3, phuong C, son tay",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-2-11:6",
                "ten_nvkt_db": "VNPT - Nguyen Van B",
                "off_duration_minutes": 19,
            },
        ],
        "Sơn Tây",
    )

    assert "👷 Nguyen Van A (2 TB)" in message
    assert "👷 Nguyen Van B (1 TB)" in message
    assert message.index("👷 Nguyen Van A (2 TB)") < message.index("[TB001] Ten TB 1 - 0912345678")
    assert message.index("👷 Nguyen Van B (1 TB)") < message.index("[TB003] Ten TB 3 - 0900000000")


def test_format_current_off_snapshot_for_doi_uses_precomputed_duration():
    message = notification_service.format_current_off_snapshot_for_doi(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB",
                "dienthoai_lh": "0912345678",
                "diachi_ld": "123 Duong Rat Dai, Phuong Trung Tam, Thi Xa Son Tay",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:1",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": "2026-04-24T08:10:00",
                "duration_minutes": 35,
            }
        ],
        "Tổ Kỹ thuật Địa bàn Sơn Tây",
    )

    assert "🚨 CẢNH BÁO THUÊ BAO OFF" in message
    assert "Kéo dài: 0,5 giờ" in message
    assert "1. [TB001] Ten TB - 0912345678" in message
    assert "OFF: 24/04/2026 08:10" in message
    assert "Đ/c: 123 Duong Rat Dai, Phuong Trung Tam, Thi" in message
    assert "Port: STY.G51_1/1" in message
    assert "Đội: Tổ Kỹ thuật Địa bàn Sơn Tây" not in message
    assert "--------" in message


def test_format_current_off_snapshot_for_doi_sorts_newest_first_and_numbers_sequentially():
    message = notification_service.format_current_off_snapshot_for_doi(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB 1",
                "dienthoai_lh": "0912345678",
                "diachi_ld": "Dia chi 1",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:1",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": "2026-04-24T08:10:00",
                "duration_minutes": 35,
            },
            {
                "ma_tb": "TB002",
                "ten_tb": "Ten TB 2",
                "dienthoai_lh": "0987654321",
                "diachi_ld": "Dia chi 2",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:2",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": "2026-04-24T08:25:00",
                "duration_minutes": 20,
            },
            {
                "ma_tb": "TB003",
                "ten_tb": "Ten TB 3",
                "dienthoai_lh": "0900000000",
                "diachi_ld": "Dia chi 3",
                "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:3",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "duration_minutes": 10,
            },
        ],
        "Tổ Kỹ thuật Địa bàn Sơn Tây",
    )

    assert message.index("1. [TB002] Ten TB 2 - 0987654321") < message.index("2. [TB001] Ten TB 1 - 0912345678")
    assert message.index("2. [TB001] Ten TB 1 - 0912345678") < message.index("3. [TB003] Ten TB 3 - 0900000000")
    assert message.count("--------") == 3


def test_format_current_off_snapshot_by_nvkt_omits_nvkt_header_and_includes_short_address():
    message = notification_service.format_current_off_snapshot_by_nvkt(
        [
            {
                "ma_tb": "TB001",
                "ten_tb": "Ten TB",
                "dienthoai_lh": "0912345678",
                "diachi_ld": "123 Duong Rat Dai, Phuong Trung Tam, Thi Xa Son Tay",
                "ten_nvkt_db": "VNPT - Nguyen Van A",
                "first_off_time": "2026-04-24T08:10:00",
                "duration_minutes": 35,
            }
        ]
    )

    assert "👷 Nguyen Van A (1 TB)" not in message
    assert "TB001 | Ten TB | 0912345678 | 08:10 | 35 phút | 123 Duong Rat Dai, Phuong Tru" in message


def test_send_current_off_snapshot_by_doi_vt_splits_messages_by_nvkt(monkeypatch):
    sent_messages = []

    async def fake_send_zalo_message_to_thread_detailed(message, thread_id, client=None):
        sent_messages.append((thread_id, message))
        return {
            "success": True,
            "thread_id": thread_id,
            "message": message,
            "error": "",
            "stdout": "",
            "stderr": "",
            "returncode": 0,
            "command": [],
        }

    monkeypatch.setattr(
        notification_service,
        "send_zalo_message_to_thread_detailed",
        fake_send_zalo_message_to_thread_detailed,
    )

    result = asyncio.run(
        notification_service.send_current_off_snapshot_by_doi_vt(
            [
                {
                    "ma_tb": "TB001",
                    "ten_tb": "Ten TB 1",
                    "dienthoai_lh": "0912345678",
                    "diachi_ld": "Dia chi 1",
                    "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:1",
                    "doi_vt": "Tổ Kỹ thuật Địa bàn Sơn Tây",
                    "ten_nvkt_db": "VNPT - Nguyen Van A",
                    "duration_minutes": 10,
                    "first_off_time": "2026-04-25T12:00:00",
                },
                {
                    "ma_tb": "TB002",
                    "ten_tb": "Ten TB 2",
                    "dienthoai_lh": "0987654321",
                    "diachi_ld": "Dia chi 2",
                    "port_id": "HNI.STY.STY.OLT.AL.2.1_1-1-1:2",
                    "doi_vt": "Tổ Kỹ thuật Địa bàn Sơn Tây",
                    "ten_nvkt_db": "VNPT - Nguyen Van B",
                    "duration_minutes": 12,
                    "first_off_time": "2026-04-25T11:58:00",
                },
            ]
        )
    )

    assert result["sent"] == 2
    assert result["failed"] == 0
    assert len(sent_messages) == 2
    assert all(thread_id == "4761925886931896176" for thread_id, _message in sent_messages)
    assert "👷 Nguyen Van A (1 TB)" in sent_messages[0][1]
    assert "👷 Nguyen Van B (1 TB)" in sent_messages[1][1]


def test_format_recovery_message_for_zalo_uses_short_olt_name_in_port(tmp_path, monkeypatch):
    mapping_file = tmp_path / "olt_mapping.xlsx"

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["OLT", "TEN_DSLAM"])
    sheet.append(["HNI.BVI.PPG.OLT.ZT.1.1", "PPG.G41"])
    workbook.save(mapping_file)

    monkeypatch.setattr(notification_service, "OLT_MAPPING_FILE", str(mapping_file))
    monkeypatch.setattr(notification_service, "_OLT_DISPLAY_NAME_CACHE", None)

    message = notification_service.format_recovery_message_for_zalo(
        {
            "ma_tb": "ongthang0369463350",
            "ten_tb": "Phùng Văn Thắng",
            "port_id": "HNI.BVI.PPG.OLT.ZT.1.1_1-1-12:14",
            "outage_duration_minutes": 592,
        }
    )

    assert "Port: PPG.G41_1-1-12:14" in message
    assert "HNI.BVI.PPG.OLT.ZT.1.1_1-1-12:14" not in message
