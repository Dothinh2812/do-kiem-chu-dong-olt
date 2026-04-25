import csv
import json
import math
import os
import re
import time
from collections import defaultdict
from datetime import datetime, time as dt_time
from typing import Dict, List, Optional

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

try:
    from .doi_vt_mapping import CANONICAL_DOI_VT_TO_THREAD, get_thread_id_for_doi_vt
    from .openzca_adapter import OpenZcaClient
except ImportError:
    from doi_vt_mapping import CANONICAL_DOI_VT_TO_THREAD, get_thread_id_for_doi_vt
    from openzca_adapter import OpenZcaClient


CONFIG_FILE = os.path.join(os.path.dirname(__file__), "config_notification.json")
OLT_MAPPING_FILE = os.path.join(os.path.dirname(__file__), "olt_mapping.xlsx")
DEFAULT_OPENZCA_PROFILE = os.environ.get("OPENZCA_PROFILE", "zalo2")
NOTIFICATION_DELIVERY_LOG_FILE = os.path.join("log_message", "notification_delivery.jsonl")
_OLT_DISPLAY_NAME_CACHE = None

load_dotenv()


def _env_bool(name: str, default: bool) -> bool:
    raw_value = os.environ.get(name)
    if raw_value is None:
        return default
    return raw_value.strip().lower() in {"1", "true", "yes", "on"}


def _env_str(name: str, default: str = "") -> str:
    return str(os.environ.get(name, default) or "").strip()


def _build_default_config() -> Dict:
    return {
        "telegram_bot_token": _env_str("TELEGRAM_BOT_TOKEN", ""),
        "telegram_chat_id": _env_str("TELEGRAM_CHAT_ID", ""),
        "enable_telegram": _env_bool("ENABLE_TELEGRAM", True),
        "enable_zalo": _env_bool("ENABLE_ZALO", True),
        "enable_individual_alert_notifications": _env_bool("ENABLE_INDIVIDUAL_ALERT_NOTIFICATIONS", True),
        "enable_wide_area_alert_notifications": _env_bool("ENABLE_WIDE_AREA_ALERT_NOTIFICATIONS", True),
        "enable_recovery_alert_notifications": _env_bool("ENABLE_RECOVERY_ALERT_NOTIFICATIONS", False),
        "individual_alert_time_window": _env_str("INDIVIDUAL_ALERT_TIME_WINDOW", ""),
        "wide_area_alert_time_window": _env_str("WIDE_AREA_ALERT_TIME_WINDOW", ""),
        "recovery_alert_time_window": _env_str("RECOVERY_ALERT_TIME_WINDOW", ""),
        "wide_area_alert_excluded_ports": _env_str("WIDE_AREA_ALERT_EXCLUDED_PORTS", ""),
    }


_ALERT_POLICY_CONFIG = {
    "outage": {
        "enabled_key": "enable_individual_alert_notifications",
        "window_key": "individual_alert_time_window",
        "label": "individual outage",
    },
    "wide_area": {
        "enabled_key": "enable_wide_area_alert_notifications",
        "window_key": "wide_area_alert_time_window",
        "label": "wide-area",
    },
    "recovery": {
        "enabled_key": "enable_recovery_alert_notifications",
        "window_key": "recovery_alert_time_window",
        "label": "recovery",
    },
}

DOI_VT_TO_THREAD = CANONICAL_DOI_VT_TO_THREAD


def load_config() -> Dict:
    config = _build_default_config()
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as handle:
                config.update(json.load(handle))
        except Exception as exc:
            print(f"⚠️ Could not load {CONFIG_FILE}: {exc}")
    return config


def _parse_hhmm(value: str) -> dt_time:
    hour_str, minute_str = value.split(":", 1)
    hour = int(hour_str)
    minute = int(minute_str)
    if hour < 0 or hour > 23 or minute < 0 or minute > 59:
        raise ValueError("hour or minute out of range")
    return dt_time(hour=hour, minute=minute)


def _is_within_time_window(window: str, now: Optional[datetime] = None) -> bool:
    normalized_window = str(window or "").strip()
    if not normalized_window:
        return True

    try:
        start_raw, end_raw = [part.strip() for part in normalized_window.split("-", 1)]
        start_time = _parse_hhmm(start_raw)
        end_time = _parse_hhmm(end_raw)
    except ValueError:
        print(f"⚠️ Invalid notification time window '{normalized_window}', allowing send.")
        return True

    current_time = (now or datetime.now()).time().replace(second=0, microsecond=0)
    if start_time <= end_time:
        return start_time <= current_time <= end_time
    return current_time >= start_time or current_time <= end_time


def get_alert_policy_status(alert_type: str, config: Optional[Dict] = None, now: Optional[datetime] = None) -> Dict[str, str]:
    policy = _ALERT_POLICY_CONFIG[alert_type]
    active_config = config or load_config()

    if not active_config.get(policy["enabled_key"], True):
        return {
            "allowed": False,
            "reason": "disabled",
            "message": f"{policy['label']} notifications disabled by config",
        }

    window = str(active_config.get(policy["window_key"], "") or "").strip()
    if not _is_within_time_window(window, now=now):
        return {
            "allowed": False,
            "reason": "outside_time_window",
            "message": f"{policy['label']} notifications outside allowed window {window}",
        }

    return {
        "allowed": True,
        "reason": "allowed",
        "message": f"{policy['label']} notifications allowed",
    }


def _normalize_excluded_port_value(value: str) -> str:
    return str(value or "").strip().upper()


def parse_wide_area_excluded_ports(raw_value: str) -> set[tuple[str, str]]:
    normalized_raw_value = str(raw_value or "").strip()
    if not normalized_raw_value:
        return set()

    entries = re.split(r"(?:\r?\n){2,}|;", normalized_raw_value)
    excluded_ports = set()

    for entry in entries:
        compact_entry = " ".join(str(entry or "").split())
        if not compact_entry:
            continue

        olt_match = re.search(r"OLT\s*:\s*([^\n,;]+?)(?=\s+PORT\s*:|$|,)", compact_entry, flags=re.IGNORECASE)
        port_match = re.search(r"PORT\s*:\s*([^\n,;]+)", compact_entry, flags=re.IGNORECASE)
        if not olt_match or not port_match:
            continue

        olt_name = _normalize_excluded_port_value(olt_match.group(1))
        port_name = _normalize_excluded_port_value(port_match.group(1))
        if olt_name and port_name:
            excluded_ports.add((olt_name, port_name))

    return excluded_ports


def is_wide_area_alert_excluded(alert, config: Optional[Dict] = None) -> bool:
    active_config = config or load_config()
    excluded_ports = parse_wide_area_excluded_ports(active_config.get("wide_area_alert_excluded_ports", ""))
    if not excluded_ports:
        return False

    normalized_olt_names = {
        _normalize_excluded_port_value(_value(alert, "olt_name", "")),
        _normalize_excluded_port_value(get_olt_display_name(_value(alert, "olt_name", ""))),
    }
    normalized_port = _normalize_excluded_port_value(_value(alert, "port", ""))

    return any((olt_name, normalized_port) in excluded_ports for olt_name in normalized_olt_names if olt_name)


def get_zalo_thread_by_doi_vt(doi_vt: str) -> Optional[str]:
    return get_thread_id_for_doi_vt(doi_vt)


def _normalize_mapping_header(value: str) -> str:
    return str(value or "").replace("\xa0", " ").strip().upper()


def _load_olt_display_name_map() -> Dict[str, str]:
    global _OLT_DISPLAY_NAME_CACHE
    if _OLT_DISPLAY_NAME_CACHE is not None:
        return _OLT_DISPLAY_NAME_CACHE

    if not os.path.exists(OLT_MAPPING_FILE):
        _OLT_DISPLAY_NAME_CACHE = {}
        return _OLT_DISPLAY_NAME_CACHE

    mapping: Dict[str, str] = {}
    try:
        workbook = load_workbook(OLT_MAPPING_FILE, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, ())
        header_index = {
            _normalize_mapping_header(header): idx for idx, header in enumerate(headers)
        }
        olt_idx = header_index.get("OLT")
        ten_dslam_idx = header_index.get("TEN_DSLAM")
        if olt_idx is None or ten_dslam_idx is None:
            _OLT_DISPLAY_NAME_CACHE = {}
            return _OLT_DISPLAY_NAME_CACHE

        for row in rows:
            olt_name = str(row[olt_idx] or "").strip()
            short_name = str(row[ten_dslam_idx] or "").strip()
            if olt_name and short_name:
                mapping[olt_name] = short_name
    except Exception as exc:
        print(f"⚠️ Could not load OLT mapping from {OLT_MAPPING_FILE}: {exc}")
        mapping = {}

    _OLT_DISPLAY_NAME_CACHE = mapping
    return _OLT_DISPLAY_NAME_CACHE


def get_olt_display_name(olt_name: str) -> str:
    normalized_name = (olt_name or "").strip()
    if not normalized_name:
        return ""
    return _load_olt_display_name_map().get(normalized_name, normalized_name)


def get_port_display_name(port_id: str) -> str:
    normalized_port_id = (port_id or "").strip()
    if not normalized_port_id:
        return ""
    olt_name, separator, remainder = normalized_port_id.partition("_")
    short_olt_name = get_olt_display_name(olt_name)
    if not separator:
        return short_olt_name
    return f"{short_olt_name}{separator}{remainder}"


def _short_nvkt(raw_value: str) -> str:
    raw_value = (raw_value or "").strip()
    if "-" in raw_value:
        return raw_value.split("-", 1)[1].strip()
    return raw_value


def _value(obj, key: str, default=None):
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


def _coerce_datetime(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    if hasattr(value, "strftime"):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
    return None


def _resolve_outage_duration_minutes(obj) -> Optional[int]:
    explicit_duration = _value(obj, "duration_minutes")
    if explicit_duration is None:
        explicit_duration = _value(obj, "off_duration_minutes")
    if explicit_duration is None:
        explicit_duration = _value(obj, "outage_duration_minutes")
    if explicit_duration not in (None, ""):
        try:
            return max(int(round(float(explicit_duration))), 0)
        except (TypeError, ValueError):
            pass

    first_off_time = _coerce_datetime(_value(obj, "first_off_time"))
    if not first_off_time:
        return None

    duration_seconds = (datetime.now() - first_off_time).total_seconds()
    if duration_seconds < 0:
        return 0
    return int(duration_seconds // 60)


def _format_outage_duration(obj) -> str:
    duration_minutes = _resolve_outage_duration_minutes(obj)
    if duration_minutes is None:
        return "-"
    return f"{duration_minutes} phút"


def _format_outage_duration_hours(obj) -> str:
    duration_minutes = _resolve_outage_duration_minutes(obj)
    if duration_minutes is None:
        return "-"
    duration_hours = math.floor((duration_minutes / 60) * 10) / 10
    return f"{duration_hours:.1f}".replace(".", ",") + " giờ"


def _truncate_address(value: str, limit: int = 40) -> str:
    return (value or "").strip()[:limit] or "-"


def _resolve_address(obj) -> str:
    return _truncate_address(_value(obj, "diachi_lapdat", "") or _value(obj, "diachi_ld", "") or "")


def _format_off_time(obj) -> str:
    first_off_time = _coerce_datetime(_value(obj, "first_off_time"))
    if first_off_time:
        return first_off_time.strftime("%d/%m/%Y %H:%M")
    raw_value = str(_value(obj, "first_off_time") or "").strip()
    return raw_value or "-"


def _format_doi_port_display_name(port_id: str) -> str:
    display_name = get_port_display_name(port_id)
    olt_name, separator, remainder = display_name.partition("_")
    if not separator:
        return display_name

    port_segments = remainder.split(":", 1)[0].split("-")
    if len(port_segments) >= 3:
        return f"{olt_name}_{port_segments[1]}/{port_segments[2]}"
    return display_name


def _sort_alerts_newest_first(alerts: List) -> List:
    def sort_key(alert):
        first_off_time = _coerce_datetime(_value(alert, "first_off_time"))
        duration_minutes = _resolve_outage_duration_minutes(alert)
        if first_off_time:
            return (0, -first_off_time.timestamp(), duration_minutes or 0, str(_value(alert, "ma_tb", "") or ""))
        return (1, duration_minutes or 0, 0, str(_value(alert, "ma_tb", "") or ""))

    return sorted(alerts, key=sort_key)


def append_notification_delivery_log(entry: Dict, log_file: str = NOTIFICATION_DELIVERY_LOG_FILE) -> None:
    os.makedirs(os.path.dirname(log_file), exist_ok=True)
    payload = {
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        **entry,
    }
    if "message_preview" not in payload:
        payload["message_preview"] = str(payload.get("message_full", "")).replace("\n", " ")[:200]
    with open(log_file, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False) + "\n")


async def send_telegram_message(
    message: str,
    bot_token: str = None,
    chat_id: str = None,
    parse_mode: str = "Markdown",
) -> bool:
    config = load_config()
    bot_token = bot_token or config.get("telegram_bot_token", "")
    chat_id = chat_id or config.get("telegram_chat_id", "")
    if not bot_token or not chat_id:
        print("⚠️ Telegram is not configured; skipping send.")
        return False

    response = requests.post(
        f"https://api.telegram.org/bot{bot_token}/sendMessage",
        json={
            "chat_id": chat_id,
            "text": message,
            "parse_mode": parse_mode,
            "disable_web_page_preview": True,
        },
        timeout=15,
    )
    if response.status_code == 200:
        return True
    print(f"❌ Telegram send failed: {response.status_code} {response.text[:200]}")
    return False


def get_openzca_client(profile: Optional[str] = None) -> OpenZcaClient:
    resolved_profile = profile if profile is not None else DEFAULT_OPENZCA_PROFILE
    return OpenZcaClient(profile=resolved_profile)


def check_openzca_auth(client: Optional[OpenZcaClient] = None) -> Dict:
    active_client = client or get_openzca_client()
    return active_client.auth_status()


async def send_zalo_message_to_thread(
    message: str,
    thread_id: str,
    client: Optional[OpenZcaClient] = None,
) -> bool:
    result = await send_zalo_message_to_thread_detailed(message, thread_id, client=client)
    return result["success"]


async def send_zalo_message_to_thread_detailed(
    message: str,
    thread_id: str,
    client: Optional[OpenZcaClient] = None,
) -> Dict:
    if not thread_id:
        return {
            "success": False,
            "thread_id": "",
            "message": message,
            "error": "missing thread_id",
        }
    try:
        active_client = client or get_openzca_client()
        if hasattr(active_client, "send_text_detailed"):
            result = active_client.send_text_detailed(thread_id, message, group=True)
            if isinstance(result, dict):
                return {
                    "success": bool(result.get("success")),
                    "thread_id": result.get("thread_id", thread_id),
                    "message": result.get("message", message),
                    "response": result.get("response", result.get("stdout", "")),
                    "error": result.get("error", ""),
                    "stdout": result.get("stdout", ""),
                    "stderr": result.get("stderr", ""),
                    "returncode": result.get("returncode", 0),
                    "command": result.get("command", []),
                }
            success = result.returncode == 0
            error_message = result.stderr or result.stdout or ""
            return {
                "success": success,
                "thread_id": thread_id,
                "message": message,
                "response": result.stdout,
                "error": "" if success else error_message,
                "stdout": result.stdout,
                "stderr": result.stderr,
                "returncode": result.returncode,
                "command": result.command,
            }
        response = active_client.send_text(thread_id, message, group=True)
    except Exception as exc:
        print(f"❌ Zalo send error via openzca: {exc}")
        return {
            "success": False,
            "thread_id": thread_id,
            "message": message,
            "error": str(exc),
            "stdout": "",
            "stderr": "",
            "returncode": -1,
            "command": [],
        }
    return {
        "success": True,
        "thread_id": thread_id,
        "message": message,
        "response": response,
        "error": "",
        "stdout": str(response or ""),
        "stderr": "",
        "returncode": 0,
        "command": [],
    }


def format_wide_area_outage_message(wide_area_alerts: List[Dict]) -> str:
    if not wide_area_alerts:
        return ""
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = [
        f"🔴 *CẢNH BÁO SỰ CỐ DIỆN RỘNG* - {now}",
        "",
        f"Phát hiện *{len(wide_area_alerts)}* port có nhiều thuê bao mất tín hiệu:",
        "",
    ]
    for idx, alert in enumerate(wide_area_alerts, 1):
        lines.append(f"*{idx}. Port {alert['port']}* - {alert['subscriber_count']} thuê bao OFF")
        lines.append(f"   🏢 OLT: {get_olt_display_name(alert['olt_name'])}")
        lines.append(f"   ⏱️ Kéo dài: {_format_outage_duration(alert)}")
        lines.append("   📋 Danh sách thuê bao:")
        for sub_idx, sub in enumerate(alert.get("subscriber_list", []), 1):
            ten_tb = (sub.get("ten_tb") or "")[:25]
            nvkt = _short_nvkt(sub.get("ten_nvkt_db") or "")
            nvkt_display = f" - {nvkt}" if nvkt else ""
            lines.append(f"      {sub_idx}. [{sub.get('ma_tb', '')}] {ten_tb}{nvkt_display}")
        lines.append("")
    return "\n".join(lines)


def format_wide_area_outage_for_zalo(wide_area_alert: Dict) -> str:
    lines = [
        "🔴 CẢNH BÁO SỰ CỐ DIỆN RỘNG",
        f"  - OLT: {get_olt_display_name(wide_area_alert['olt_name'])}",
        f"  - Port: {wide_area_alert['port']}",
        f"  - Số thuê bao OFF: {wide_area_alert['subscriber_count']}",
        f"  - Kéo dài: {_format_outage_duration(wide_area_alert)}",
        "  - Danh sách thuê bao:",
    ]
    for idx, sub in enumerate(wide_area_alert.get("subscriber_list", []), 1):
        ten_tb = (sub.get("ten_tb") or "")[:25]
        nvkt = _short_nvkt(sub.get("ten_nvkt_db") or "")
        nvkt_display = f" - {nvkt}" if nvkt else ""
        lines.append(f"      {idx}. [{sub.get('ma_tb', '')}] {ten_tb}{nvkt_display}")
    return "\n".join(lines)


def format_consolidated_outage_by_nvkt(alerts: List, for_zalo: bool = False) -> str:
    groups = defaultdict(list)
    for alert in alerts:
        nvkt = _short_nvkt(_value(alert, "ten_nvkt_db", "") or "")
        groups[nvkt or "Chưa gán NVKT"].append(alert)

    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = [
        f"🚨 {'THUÊ BAO OFF' if for_zalo else '*THUÊ BAO OFF*'} - {now}",
        "",
    ]
    for nvkt, items in groups.items():
        lines.append(f"👷 {nvkt} ({len(items)} TB)")
        for alert in items:
            ma_tb = _value(alert, "ma_tb", "") or ""
            ten_tb = _value(alert, "ten_tb", "") or ""
            sdt = _value(alert, "dienthoai_lh", "") or "-"
            off_time = _coerce_datetime(_value(alert, "first_off_time"))
            off_time_str = off_time.strftime("%H:%M") if off_time else str(_value(alert, "first_off_time") or "-")
            lines.append(f"  {ma_tb} | {ten_tb} | {sdt} | {off_time_str} | {_format_outage_duration(alert)}")
        lines.append("")
    return "\n".join(lines)


def format_consolidated_outage_for_doi(alerts: List, doi_vt: str) -> str:
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = [
        f"🚨 CẢNH BÁO THUÊ BAO OFF - {now}",
        "",
    ]
    groups = defaultdict(list)
    for alert in alerts:
        nvkt = _short_nvkt(_value(alert, "ten_nvkt_db", "") or "")
        groups[nvkt or "Chưa gán NVKT"].append(alert)

    alert_index = 1
    for nvkt, items in groups.items():
        lines.append(f"👷 {nvkt} ({len(items)} TB)")
        for alert in _sort_alerts_newest_first(items):
            ma_tb = _value(alert, "ma_tb", "") or ""
            ten_tb = _value(alert, "ten_tb", "") or ""
            sdt = _value(alert, "dienthoai_lh", "") or "-"
            dia_chi = _resolve_address(alert)
            port_id = _value(alert, "port_id", "") or ""
            lines.append(f"{alert_index}. [{ma_tb}] {ten_tb} - {sdt}")
            lines.append(f"   OFF: {_format_off_time(alert)}")
            lines.append(f"   Đ/c: {dia_chi}")
            lines.append(f"   Port: {_format_doi_port_display_name(port_id)}")
            lines.append(f"   Kéo dài: {_format_outage_duration_hours(alert)}")
            lines.append("--------")
            alert_index += 1
        lines.append("")
    return "\n".join(lines)


def format_recovery_message(alerts: List) -> str:
    if not alerts:
        return ""
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = [
        f"✅ *THÔNG BÁO PHỤC HỒI* - {now}",
        "",
        f"*{len(alerts)}* thuê bao đã phục hồi kết nối:",
        "",
    ]
    for idx, alert in enumerate(alerts, 1):
        ma_tb = _value(alert, "ma_tb", "") or ""
        ten_tb = _value(alert, "ten_tb", "") or ""
        port_id = _value(alert, "port_id", "") or ""
        olt_name = _value(alert, "olt_name", "") or ""
        duration = _value(alert, "outage_duration_minutes", 0) or 0
        lines.append(f"*{idx}. [{ma_tb}] {ten_tb}*")
        lines.append(f"   📍 Port: `{get_port_display_name(port_id)}`")
        lines.append(f"   🏢 OLT: {get_olt_display_name(olt_name)}")
        lines.append(f"   ⏱️ Thời gian mất: {duration} phút")
        lines.append("")
    return "\n".join(lines)


def format_recovery_message_for_zalo(alert) -> str:
    ma_tb = _value(alert, "ma_tb", "") or ""
    ten_tb = _value(alert, "ten_tb", "") or ""
    port_id = _value(alert, "port_id", "") or ""
    duration = _value(alert, "outage_duration_minutes", 0) or 0
    return "\n".join(
        [
            "✅ THÔNG BÁO PHỤC HỒI",
            f"  - Mã TB: {ma_tb}",
            f"  - Tên TB: {ten_tb}",
            f"  - Port: {get_port_display_name(port_id)}",
            f"  - Thời gian mất: {duration} phút",
        ]
    )


async def send_consolidated_outage_telegram(alerts: List) -> bool:
    message = format_consolidated_outage_by_nvkt(alerts, for_zalo=False)
    if not message:
        return True
    return await send_telegram_message(message)


async def send_consolidated_outage_by_doi_vt(alerts: List) -> Dict[str, int]:
    results = {"sent": 0, "failed": 0, "no_thread": 0, "deliveries": []}
    groups = defaultdict(list)
    for alert in alerts:
        doi_vt = _value(alert, "doi_vt", "") or "Không xác định"
        groups[doi_vt].append(alert)

    log_entries = []
    for doi_vt, items in groups.items():
        thread_id = get_zalo_thread_by_doi_vt(doi_vt)
        message = format_consolidated_outage_for_doi(items, doi_vt)
        alert_ids = [_value(alert, "id") for alert in items if _value(alert, "id")]
        if not thread_id:
            results["no_thread"] += 1
            log_entries.append(
                {
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "alert_type": "consolidated_outage",
                    "doi_vt": doi_vt,
                    "alert_count": len(items),
                    "thread_id": "N/A",
                    "status": "NO_THREAD",
                    "error": "",
                }
            )
            results["deliveries"].append(
                {
                    "channel": "zalo",
                    "alert_type": "outage",
                    "doi_vt": doi_vt,
                    "thread_id": "",
                    "status": "NO_THREAD",
                    "message_full": message,
                    "alert_ids": alert_ids,
                    "alert_count": len(items),
                    "error": "DOI_VT not mapped",
                    "stdout": "",
                    "stderr": "",
                    "returncode": None,
                    "command": [],
                }
            )
            continue
        delivery_result = await send_zalo_message_to_thread_detailed(message, thread_id)
        success = delivery_result["success"]
        results["sent" if success else "failed"] += 1
        log_entries.append(
            {
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "alert_type": "consolidated_outage",
                "doi_vt": doi_vt,
                "alert_count": len(items),
                "thread_id": thread_id,
                "status": "SUCCESS" if success else "FAILED",
                "error": delivery_result.get("error", ""),
            }
        )
        results["deliveries"].append(
            {
                "channel": "zalo",
                "alert_type": "outage",
                "doi_vt": doi_vt,
                "thread_id": thread_id,
                "status": "SUCCESS" if success else "FAILED",
                "message_full": message,
                "alert_ids": alert_ids,
                "alert_count": len(items),
                "error": delivery_result.get("error", ""),
                "stdout": delivery_result.get("stdout", ""),
                "stderr": delivery_result.get("stderr", ""),
                "returncode": delivery_result.get("returncode"),
                "command": delivery_result.get("command", []),
            }
        )
        time.sleep(0.5)

    if log_entries:
        log_dir = "log_message"
        os.makedirs(log_dir, exist_ok=True)
        log_file = os.path.join(log_dir, "zalo_consolidated_outage_log.csv")
        file_exists = os.path.exists(log_file)
        with open(log_file, "a", newline="", encoding="utf-8-sig") as handle:
            fieldnames = ["timestamp", "alert_type", "doi_vt", "alert_count", "thread_id", "status", "error"]
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            if not file_exists:
                writer.writeheader()
            writer.writerows(log_entries)

    return results


def format_current_off_snapshot_by_nvkt(alerts: List, for_zalo: bool = False) -> str:
    return format_consolidated_outage_by_nvkt(alerts, for_zalo=for_zalo)


def format_current_off_snapshot_for_doi(alerts: List, doi_vt: str) -> str:
    return format_consolidated_outage_for_doi(alerts, doi_vt)


async def send_current_off_snapshot_telegram(alerts: List) -> bool:
    message = format_current_off_snapshot_by_nvkt(alerts, for_zalo=False)
    if not message:
        return True
    return await send_telegram_message(message)


async def send_current_off_snapshot_by_doi_vt(alerts: List) -> Dict[str, int]:
    results = {"sent": 0, "failed": 0, "no_thread": 0, "deliveries": []}
    groups = defaultdict(list)
    for alert in alerts:
        doi_vt = _value(alert, "doi_vt", "") or "Không xác định"
        groups[doi_vt].append(alert)

    log_entries = []
    for doi_vt, items in groups.items():
        thread_id = get_zalo_thread_by_doi_vt(doi_vt)
        nvkt_groups = defaultdict(list)
        for alert in items:
            nvkt = _short_nvkt(_value(alert, "ten_nvkt_db", "") or "")
            nvkt_groups[nvkt or "Chưa gán NVKT"].append(alert)

        for nvkt_items in nvkt_groups.values():
            message = format_current_off_snapshot_for_doi(nvkt_items, doi_vt)
            if not thread_id:
                results["no_thread"] += 1
                log_entries.append(
                    {
                        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "alert_type": "current_off_snapshot",
                        "doi_vt": doi_vt,
                        "alert_count": len(nvkt_items),
                        "thread_id": "N/A",
                        "status": "NO_THREAD",
                        "error": "",
                    }
                )
                results["deliveries"].append(
                    {
                        "channel": "zalo",
                        "alert_type": "outage",
                        "doi_vt": doi_vt,
                        "thread_id": "",
                        "status": "NO_THREAD",
                        "message_full": message,
                        "alert_ids": [],
                        "alert_count": len(nvkt_items),
                        "error": "DOI_VT not mapped",
                        "stdout": "",
                        "stderr": "",
                        "returncode": None,
                        "command": [],
                    }
                )
                continue
            delivery_result = await send_zalo_message_to_thread_detailed(message, thread_id)
            success = delivery_result["success"]
            results["sent" if success else "failed"] += 1
            log_entries.append(
                {
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "alert_type": "current_off_snapshot",
                    "doi_vt": doi_vt,
                    "alert_count": len(nvkt_items),
                    "thread_id": thread_id,
                    "status": "SUCCESS" if success else "FAILED",
                    "error": delivery_result.get("error", ""),
                }
            )
            results["deliveries"].append(
                {
                    "channel": "zalo",
                    "alert_type": "outage",
                    "doi_vt": doi_vt,
                    "thread_id": thread_id,
                    "status": "SUCCESS" if success else "FAILED",
                    "message_full": message,
                    "alert_ids": [],
                    "alert_count": len(nvkt_items),
                    "error": delivery_result.get("error", ""),
                    "stdout": delivery_result.get("stdout", ""),
                    "stderr": delivery_result.get("stderr", ""),
                    "returncode": delivery_result.get("returncode"),
                    "command": delivery_result.get("command", []),
                }
            )
            time.sleep(0.5)

    if log_entries:
        log_dir = "log_message"
        os.makedirs(log_dir, exist_ok=True)
        log_file = os.path.join(log_dir, "zalo_consolidated_outage_log.csv")
        file_exists = os.path.exists(log_file)
        with open(log_file, "a", newline="", encoding="utf-8-sig") as handle:
            fieldnames = ["timestamp", "alert_type", "doi_vt", "alert_count", "thread_id", "status", "error"]
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            if not file_exists:
                writer.writeheader()
            writer.writerows(log_entries)

    return results


async def send_recovery_alert_telegram(alerts: List) -> bool:
    message = format_recovery_message(alerts)
    if not message:
        return True
    return await send_telegram_message(message)


async def send_recovery_alerts_by_doi_vt(alerts: List) -> Dict[str, int]:
    results = {"sent": 0, "failed": 0, "no_thread": 0, "deliveries": []}
    for alert in alerts:
        doi_vt = _value(alert, "doi_vt", "") or ""
        thread_id = get_zalo_thread_by_doi_vt(doi_vt)
        message = format_recovery_message_for_zalo(alert)
        alert_id = _value(alert, "id")
        if not thread_id:
            results["no_thread"] += 1
            results["deliveries"].append(
                {
                    "channel": "zalo",
                    "alert_type": "recovery",
                    "doi_vt": doi_vt,
                    "thread_id": "",
                    "status": "NO_THREAD",
                    "message_full": message,
                    "alert_ids": [alert_id] if alert_id else [],
                    "alert_count": 1,
                    "error": "DOI_VT not mapped",
                    "stdout": "",
                    "stderr": "",
                    "returncode": None,
                    "command": [],
                }
            )
            continue
        delivery_result = await send_zalo_message_to_thread_detailed(message, thread_id)
        success = delivery_result["success"]
        results["sent" if success else "failed"] += 1
        results["deliveries"].append(
            {
                "channel": "zalo",
                "alert_type": "recovery",
                "doi_vt": doi_vt,
                "thread_id": thread_id,
                "status": "SUCCESS" if success else "FAILED",
                "message_full": message,
                "alert_ids": [alert_id] if alert_id else [],
                "alert_count": 1,
                "error": delivery_result.get("error", ""),
                "stdout": delivery_result.get("stdout", ""),
                "stderr": delivery_result.get("stderr", ""),
                "returncode": delivery_result.get("returncode"),
                "command": delivery_result.get("command", []),
            }
        )
        time.sleep(0.5)
    return results
