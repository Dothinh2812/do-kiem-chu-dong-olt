#!/usr/bin/env python3
"""Tra Zalo ID từ danh sách nhân viên (dsnv.xlsx) bằng openzca friend find."""

import json
import subprocess
import time
from pathlib import Path

import openpyxl

OPENZCA_BIN = "/home/vtst/.nvm/versions/node/v22.22.2/bin/openzca"
NODE_BIN = "/home/vtst/.nvm/versions/node/v22.22.2/bin/node"
PROFILE = "TTVTST"
INPUT_FILE = Path(__file__).parent / "dsnv.xlsx"
OUTPUT_FILE = Path(__file__).parent / "dsnv_zalo.xlsx"


def find_zalo_by_phone(phone: str) -> dict | None:
    """Gọi openzca friend find, trả về dict chứa userId hoặc None."""
    cmd = [NODE_BIN, OPENZCA_BIN, "--profile", PROFILE, "friend", "find", "--json", phone]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, check=False, timeout=30)
        if proc.returncode == 0 and proc.stdout.strip():
            users = json.loads(proc.stdout)
            if users:
                u = users[0]
                return {
                    "userId": u.get("userId"),
                    "zalo_name": u.get("display_name") or u.get("zalo_name"),
                    "globalId": u.get("globalId"),
                }
    except Exception:
        pass
    return None


def main():
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    headers.extend(["zalo_userId", "zalo_name", "zalo_globalId"])
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.append(headers)

    total = ws.max_row - 1
    found = 0
    failed = 0

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        phone = row[4]  # cột sdt (index 4)
        name = row[2]   # cột NVKT (tên)
        row_data = list(row)

        print(f"[{i}/{total}] {name} ({phone}) ... ", end="", flush=True)

        if not phone:
            row_data.extend(["", "", ""])
            print("SKIP (no phone)")
        else:
            phone_str = str(phone).strip()
            result = find_zalo_by_phone(phone_str)
            if result:
                row_data.extend([
                    result["userId"],
                    result["zalo_name"],
                    result["globalId"],
                ])
                print(f"OK -> {result['userId']}")
                found += 1
            else:
                row_data.extend(["", "", ""])
                print("NOT FOUND")
                failed += 1

        out_ws.append(row_data)

        # Rate limit: 0.5s giữa mỗi request
        if i < total:
            time.sleep(0.5)

    out_wb.save(OUTPUT_FILE)
    print(f"\n=== DONE ===")
    print(f"Tổng: {total} | Tìm thấy: {found} | Không tìm thấy: {failed}")
    print(f"Kết quả lưu tại: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
